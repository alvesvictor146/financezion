import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\planilha\Victor Orçamento Anual.xlsx", data_only=True)
ws_cartao = wb['CARTAO AGOSTO']
cartao_sums = {}
for r in range(4, 150):
    cat = ws_cartao.cell(r, 2).value
    val = ws_cartao.cell(r, 3).value
    if cat and val is not None:
        cat = str(cat).strip()
        cartao_sums[cat] = cartao_sums.get(cat, 0.0) + float(val)

print("CARTAO AGOSTO SUMS BY 'O que?':")
for k, v in sorted(cartao_sums.items(), key=lambda x: -x[1]):
    print(f"{k:45s}: R$ {v:10.2f}")

ws26 = wb['Orçamento 2026']
print("\nCOMPARED TO ORÇAMENTO 2026 (Agosto - Col D):")
for r in range(17, 52):
    cat_orc = ws26.cell(r, 2).value
    val_orc = ws26.cell(r, 4).value
    if cat_orc:
        cat_orc = str(cat_orc).strip()
        c_sum = cartao_sums.get(cat_orc, 0.0)
        print(f"{cat_orc:45s} | Cartão: R$ {c_sum:8.2f} | Orçamento D{r}: R$ {float(val_orc or 0):8.2f}")
