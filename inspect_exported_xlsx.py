import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\Victor_Orcamento_Anual_Atualizado_Teste.xlsx", data_only=True)
print("SHEETS IN EXPORTED XLSX:")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"Sheet '{name}': {ws.max_row} rows, {ws.max_column} cols")

ws26 = wb['Orçamento 2026']
print("\n--- AMOSTRA DO ORÇAMENTO 2026 EXPORTADO ---")
for r in [2, 3, 4, 15, 17, 21, 22, 34, 36, 45, 51, 57, 59]:
    row_vals = [f"Col {openpyxl.utils.get_column_letter(c)}: {ws26.cell(r, c).value}" for c in range(2, 10) if ws26.cell(r, c).value is not None]
    print(f"R{r:02d}: " + " | ".join(row_vals))

ws_cartao = wb['CARTAO AGOSTO']
print("\n--- AMOSTRA DA ABA CARTAO AGOSTO EXPORTADA ---")
for r in range(2, 8):
    row_vals = [f"Col {openpyxl.utils.get_column_letter(c)}: {ws_cartao.cell(r, c).value}" for c in range(2, 9) if ws_cartao.cell(r, c).value is not None]
    print(f"R{r:02d}: " + " | ".join(row_vals))

last_row = ws_cartao.max_row
print(f"Total row (R{last_row}):", [ws_cartao.cell(last_row, c).value for c in range(2, 8)])
