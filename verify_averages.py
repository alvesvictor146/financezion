import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\planilha\Victor Orçamento Anual.xlsx", data_only=True)
ws26 = wb['Orçamento 2026']

months = ['Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
cols = ['C', 'D', 'E', 'F', 'G', 'H']

print("=== VERIFICAÇÃO DETALHADA ORÇAMENTO 2026 ===")
for m, c in zip(months, cols):
    col_idx = openpyxl.utils.column_index_from_string(c)
    rec = ws26.cell(15, col_idx).value
    ess = ws26.cell(36, col_idx).value
    est = ws26.cell(51, col_idx).value
    inv = ws26.cell(57, col_idx).value
    sobra = ws26.cell(59, col_idx).value
    gastos = (ess or 0) + (est or 0)
    print(f"{m:10s} ({c}): Receita={rec:10.2f} | Essencial={ess:10.2f} | Estilo={est:10.2f} | Gastos Tot={gastos:10.2f} | Invest={inv:8.2f} | Sobra={sobra:10.2f}")

# Meses futuros: Setembro, Outubro, Novembro, Dezembro (cols E, F, G, H)
fut_recs = [ws26.cell(15, openpyxl.utils.column_index_from_string(c)).value for c in ['E', 'F', 'G', 'H']]
fut_ess = [ws26.cell(36, openpyxl.utils.column_index_from_string(c)).value for c in ['E', 'F', 'G', 'H']]
fut_est = [ws26.cell(51, openpyxl.utils.column_index_from_string(c)).value for c in ['E', 'F', 'G', 'H']]
fut_gastos = [e + s for e, s in zip(fut_ess, fut_est)]
fut_inv = [ws26.cell(57, openpyxl.utils.column_index_from_string(c)).value for c in ['E', 'F', 'G', 'H']]
fut_sobra = [ws26.cell(59, openpyxl.utils.column_index_from_string(c)).value for c in ['E', 'F', 'G', 'H']]

print("\n=== MÉDIAS DOS MESES FUTUROS (SETEMBRO A DEZEMBRO) ===")
print(f"Média Receita Futura : R$ {sum(fut_recs)/4:10.2f} / mês")
print(f"Média Gastos Futuros : R$ {sum(fut_gastos)/4:10.2f} / mês (Fixas: R$ {sum(fut_ess)/4:10.2f} + Estilo: R$ {sum(fut_est)/4:10.2f})")
print(f"Média Invest. Futuro : R$ {sum(fut_inv)/4:10.2f} / mês")
print(f"Média Sobra Futura   : R$ {sum(fut_sobra)/4:10.2f} / mês")
print(f"Total Sobra Futura (Set-Dez): R$ {sum(fut_sobra):10.2f}")
print(f"Total Sobra D59:H59 (I59 na planilha): R$ {ws26.cell(59, 9).value:10.2f}")
