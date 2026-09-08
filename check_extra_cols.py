import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\planilha\Victor Orçamento Anual.xlsx", data_only=True)
ws = wb['Orçamento 2026']

print("=== MONTHS IN ORÇAMENTO 2026 ===")
for c in range(3, 10):
    col = openpyxl.utils.get_column_letter(c)
    nome = ws.cell(3, c).value
    rec = ws.cell(15, c).value
    ess = ws.cell(36, c).value
    est = ws.cell(51, c).value
    inv = ws.cell(57, c).value
    sobra = ws.cell(59, c).value
    print(f"Col {col} ({nome}): Receita={rec} | Essencial={ess} | Estilo={est} | Invest={inv} | Sobra={sobra}")
