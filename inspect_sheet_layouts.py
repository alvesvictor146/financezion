import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\planilha\Victor Orçamento Anual.xlsx", data_only=True)

print("=== WORKBOOK SHEETS ===")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- Sheet: {name} (Max rows: {ws.max_row}, Max cols: {ws.max_column}) ---")
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or '') for c in range(1, min(10, ws.max_column + 1))]
        if any(row_vals):
            print(f"R{r:02d}: " + " | ".join(row_vals))
