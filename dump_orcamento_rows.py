import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\planilha\Victor Orçamento Anual.xlsx", data_only=False)

ws26 = wb['Orçamento 2026']
print("ROWS IN ORÇAMENTO 2026:")
categories_map = []
for r in range(3, 62):
    cat = ws26.cell(r, 2).value
    c3 = ws26.cell(r, 3).value
    d3 = ws26.cell(r, 4).value
    e3 = ws26.cell(r, 5).value
    f3 = ws26.cell(r, 6).value
    g3 = ws26.cell(r, 7).value
    h3 = ws26.cell(r, 8).value
    if cat is not None:
        categories_map.append({
            'row': r,
            'label': cat,
            'julho': c3,
            'agosto': d3,
            'setembro': e3,
            'outubro': f3,
            'novembro': g3,
            'dezembro': h3
        })
        print(f"R{r:02d}: {str(cat):40s} | Jul={c3} | Ago={d3} | Set={e3} | Out={f3} | Nov={g3} | Dez={h3}")

import json
with open(r"C:\Users\user0636\.gemini\antigravity-ide\scratch\financeiro\orcamento_2026_rows.json", "w", encoding="utf-8") as f:
    json.dump(categories_map, f, ensure_ascii=False, indent=2)
print("\nSalvo orcamento_2026_rows.json com sucesso!")
